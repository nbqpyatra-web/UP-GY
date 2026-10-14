#!/usr/bin/env python3
"""Build the data snapshot for the UP Gunvatta Yatra portal.

Combines four sources and writes them into portal/index.html (in the portal-data
script block) so the portal is a single self-contained file:

  1. data/Gunvatta_Yatra_UP_Route_Plan.xlsx - the 4-team route plan: 63 Friday yatras,
                                              10 reserve slots, holidays, weekly rhythm,
                                              notes & risks
  2. data/Planned_Yatras_All_Phases.xlsx    - hospital figures per district ((VF) Final
                                              Sheet copy: PM-JAY, NABH, gap, baseline)
  3. Freshworks CRM (live)                  - registrations, attendance and lead status
                                              per event
  4. portal/up_districts.geojson            - simplified district boundaries

Only aggregate counts leave the CRM; no contact names, emails or phone numbers.

Usage:  python3 portal/build_data.py            (needs Freshworks credentials,
                                                  see freshworks_mcp/README.md)
        python3 portal/build_data.py --cached   (reuse the last CRM pull, no API calls)
        python3 portal/build_data.py --no-crm   (planner only, for offline work)
"""

import datetime as dt
import json
import os
import re
import sys
import time

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "freshworks_mcp"))
import server as fw  # noqa: E402

ROUTE_PLAN = os.path.join(ROOT, "data", "Gunvatta_Yatra_UP_Route_Plan.xlsx")
PLANNER = os.path.join(ROOT, "data", "Planned_Yatras_All_Phases.xlsx")  # hospital figures only
GEOJSON = os.path.join(ROOT, "portal", "up_districts.geojson")
INDEX = os.path.join(ROOT, "portal", "index.html")
CACHE = os.path.join(ROOT, "portal", ".cache", "crm_events.json")

CIRCULAR_DATE = "2026-05-15"
PROGRAMME_START = "2026-10-23"   # UP GY launch (Lucknow, joint); earlier events are history
MATCH_DAYS = 3                   # a CRM event within this many days of the plan date matches
IST = dt.timezone(dt.timedelta(hours=5, minutes=30))

# CRM filters for "UP Gunvatta Yatra" events. State is matched on either state field.
CRM_BASE = [
    {"attribute": "cf_boarddevision", "operator": "is_in", "value": ["NABH"]},
    {"attribute": "cf_event_type_A", "operator": "is_in",
     "value": ["Yatra", "Gunvatta Mohotsav", "Gunvatta Yatra"]},
]
STATE_RULES = [
    {"attribute": "cf_yatra_state", "operator": "is_in", "value": ["Uttar Pradesh"]},
    {"attribute": "cf_event_state_new_dropdown", "operator": "is_in", "value": ["Uttar Pradesh"]},
]
EXCLUDED_EVENTS = {"NABH_ADDITIONAL_Y"}  # undated catch-all bucket, not a real event

# Contact Status (contact_status_id) -> portal key. Filters need the numeric id.
STATUSES = {
    "hot": ("Hot", 402001708736),
    "warm": ("Warm", 402001708737),
    "aip": ("Application in progress", 402001706085),
    "submitted": ("Application submitted", 402001706086),
    "won": ("Won", 402001578066),
    "cold": ("Cold", 402001708738),
    "engaged": ("Engaged", 402001709856),
    "lost": ("Lost", 402001578067),
    "existing": ("Existing", 402001723401),
    "registered": ("Registered", 402001711168),
    "qualified": ("Qualified", 402001728215),
    "unqualified": ("Unqualified", 402001710291),
    "attended_status": ("Attended", 402001711169),
    "not_attended_status": ("Not Attended", 402001711170),
    "new": ("New", 402001709857),
    "mitra": ("Mitra Agent", 402001698263),
    "self_guided": ("Self - Guided", 402001706080),
}
ATTENDANCE = ["Attended", "Not Attended", "Registered", "Postponed", "Cancelled"]

# CRM city / GeoJSON spelling -> planner district name.
DISTRICT_ALIASES = {
    "Kanpur": "Kanpur Nagar", "Allahabad": "Prayagraj", "Noida": "Gautam Buddha Nagar",
    "Greater Noida": "Gautam Buddha Nagar", "Dadri": "Gautam Buddha Nagar",
    "Firozabadad": "Firozabad", "Sahibabad": "Ghaziabad", "Modinagar": "Ghaziabad",
    "Mughalsarai": "Chandauli", "Faizabad": "Ayodhya", "Bara Banki": "Barabanki",
    "Bulandshahr": "Bulandshahar", "Kheri": "Lakhimpur Kheri", "Kushinagar": "Kushi Nagar",
    "Mahrajganj": "Maharajganj", "Shrawasti": "Shravasti", "Siddharthnagar": "Siddharth Nagar",
    "Raebareli": "Rae Bareli", "Orai": "Jalaun", "Sant Ravidas Nagar": "Bhadohi",
}


def canon(name):
    name = " ".join(str(name or "").split())
    return DISTRICT_ALIASES.get(name, name)


# --------------------------------------------------------------------------- plan

def iso(v):
    return v.date().isoformat() if hasattr(v, "date") else v


def sheet_rows(ws, first_header):
    """Rows below the header row whose first cell is `first_header`, as dicts."""
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == first_header)
    head = rows[hi]
    out = []
    for r in rows[hi + 1:]:
        if not any(v is not None for v in r):
            continue
        out.append({h: iso(v) for h, v in zip(head, r) if h})
    return out


def load_districts():
    """Hospital figures per district, from the (VF) Final Sheet copy in the earlier planner."""
    wb = openpyxl.load_workbook(PLANNER, data_only=True)
    districts = {}
    for g in sheet_rows(wb["All Yatras"], "Phase"):
        if g["Phase"] == "TOTAL":
            continue
        name = canon(g["District"])
        districts[name] = {
            "name": name, "mandal": g["Mandal"], "zone": g["Zone"], "tier": g["Tier"],
            "universe": g["PM-JAY Universe"] or 0, "nabh": g["NABH Total"] or 0,
            "cert": g["Certified"] or 0, "accr": g["Accredited"] or 0,
            "gap": max(0, g["Potential Gap"] or 0),  # floored at 0 (project rule)
            "gapRaw": g["Potential Gap"] or 0,
            "baseline": {  # (VF) Final Sheet figures as of 06 Oct 2026
                "yatras": g["Yatras Conducted"] or 0, "attended": g["Attended"] or 0,
                "won": g["Won"] or 0, "aip": g["Application in Progress"] or 0,
                "hot": g["Hot"] or 0, "warm": g["Warm"] or 0,
            },
        }
    return districts


TEAM_KEYS = {"Joint (All 4)": "J", "Team 1 – West": "T1", "Team 2 – East": "T2",
             "Team 3 – Central": "T3", "Team 4 – South-West & Bundelkhand": "T4"}


def blank(v):
    return None if v in (None, "", "—", "-") else v


def load_route():
    wb = openpyxl.load_workbook(ROUTE_PLAN, data_only=True)
    yatras, reserves = [], []
    for r in sheet_rows(wb["Route Plan"], "S.No"):
        if not isinstance(r["S.No"], int):
            continue
        team = TEAM_KEYS[r["Team"]]
        leg = r["Leg"]
        base = {
            "id": leg if team == "J" else f"{team}-{leg}", "sno": r["S.No"], "team": team, "teamName": r["Team"],
            "teamZone": r["Zone"], "circuit": r["Phase / Circuit"], "leg": leg,
            "date": r["Yatra Date"], "day": r["Day"], "onGround": r["Team on ground by"],
            "calendar": r["Calendar check"], "remarks": blank(r["Remarks"]),
        }
        if r["Status"] == "Reserve":
            reserves.append(base)
            continue
        base.update({
            "mandal": r["Mandal"], "host": canon(r["District / Cluster (Host)"].replace(" Cluster", "")),
            "label": r["District / Cluster (Host)"],
            "districts": [canon(d) for d in str(r["Districts covered"]).split(",")],
            "travelFrom": blank(r["Travel from"]), "km": blank(r["Approx. road km"]),
            "earlier": blank(r["Earlier GY visit"]), "venue": blank(r["Venue (Primary)"]),
            "fallback": blank(r["Fallback"]), "refreshAV": blank(r["Refreshment & AV"]),
            "faculty": blank(r["NABH Faculty"]), "spoc": blank(r["District SPOC (CMO office)"]),
            "turnout": blank(r["Turnout target"]), "planStatus": r["Status"],
        })
        base["cluster"] = r["District / Cluster (Host)"] if len(base["districts"]) > 1 else None
        yatras.append(base)

    teams = []
    for r in sheet_rows(wb["Summary"], "Team"):
        if r["Team"] in TEAM_KEYS:
            teams.append({"key": TEAM_KEYS[r["Team"]], "name": r["Team"], "zone": r["Zone"],
                          "member": blank(r["Team member (1 person)"]), "base": blank(r["Home base"])})
        if r["Team"] == "TOTAL":
            break

    holidays = [{"date": r["Date"], "day": r["Day"], "name": r["Holiday / Festival"],
                 "type": r["Type"], "blocks": r["Blocks Yatra?"], "note": blank(r["Note"])}
                for r in sheet_rows(wb["UP Holidays"], "Date")]
    fridays = [{"friday": r["Friday"], "status": r["FINAL STATUS"], "date": r["Yatra date to use"],
                "count": r["Yatras scheduled"], "override": blank(r["Planning override"])}
               for r in sheet_rows(wb["Friday Calendar"], "Friday") if r["Friday"] != "TOTAL"]
    rhythm = [{"when": r["When"], "activity": r["Activity"], "output": r["Output / check"], "with": r["With"]}
              for r in sheet_rows(wb["Weekly Rhythm"], "When")]
    risks = [{"no": r["#"], "item": r["Item"], "detail": r["Detail / Action"], "level": r["Level"]}
             for r in sheet_rows(wb["Notes & Risks"], "#")]
    rules = [str(r[0]) for r in wb["Summary"].iter_rows(values_only=True)
             if r[0] and re.match(r"^\d+\. ", str(r[0]))]
    return {"yatras": yatras, "reserves": reserves, "teams": teams, "holidays": holidays,
            "fridays": fridays, "rhythm": rhythm, "risks": risks, "rules": rules}


# --------------------------------------------------------------------------- CRM

class Budget:
    """Keeps calls under the Freshworks per-minute limit (400/min)."""

    def __init__(self, per_min=330):
        self.per_min, self.stamps, self.total = per_min, [], 0

    def tick(self):
        now = time.time()
        self.stamps = [t for t in self.stamps if now - t < 60]
        if len(self.stamps) >= self.per_min:
            time.sleep(60 - (now - self.stamps[0]) + 0.5)
        self.stamps.append(time.time())
        self.total += 1


budget = Budget()


def call(method, path, body=None):
    for attempt in range(4):
        budget.tick()
        try:
            return fw.api(method, path, body=body)
        except fw.FreshworksError as e:
            transient = "Could not reach" in str(e) or "HTTP 429" in str(e) or "HTTP 50" in str(e)
            if not transient or attempt == 3:
                raise
            time.sleep(2 ** (attempt + 1))


def search(rules, per_page=1):
    return call("POST", "filtered_search/contact",
                {"filter_rule": rules, "page": 1, "per_page": per_page})


def count(rules):
    return search(rules)["meta"]["total"]


def discover_events():
    """Find every distinct event name: take one contact, read its event, exclude, repeat."""
    events = {}
    for state in STATE_RULES:
        while True:
            rules = CRM_BASE + [state]
            known = list(events) + list(EXCLUDED_EVENTS)
            rules.append({"attribute": "cf_event_type", "operator": "is_not_in", "value": known})
            res = search(rules)
            if not res["contacts"]:
                break
            cf = call("GET", f"contacts/{res['contacts'][0]['id']}")["contact"]["custom_field"]
            name = cf.get("cf_event_type")
            if not name or name in events:
                print(f"  stopping: contact without a usable event name ({name!r})", file=sys.stderr)
                break
            events[name] = cf
            print(f"  found {name} ({res['meta']['total']} contacts left)", file=sys.stderr)
    return events


def partner_model(partner):
    p = (partner or "").lower()
    if any(k in p for k in ("runner", "standalone")):
        return "direct"
    return "partner" if p else "unknown"


def crm_events():
    raw = discover_events()
    # Statuses in use anywhere in UP, so per-event queries skip empty ones.
    any_state = CRM_BASE + [STATE_RULES[0]]
    used = {k: v for k, v in STATUSES.items()
            if count(any_state + [{"attribute": "contact_status_id", "operator": "is_in",
                                   "value": [str(v[1])]}]) or
            count(CRM_BASE + [STATE_RULES[1], {"attribute": "contact_status_id",
                                              "operator": "is_in", "value": [str(v[1])]}])}
    att_used = [a for a in ATTENDANCE
                if count(any_state + [{"attribute": "cf_attendance", "operator": "is_in", "value": [a]}])
                or count(CRM_BASE + [STATE_RULES[1], {"attribute": "cf_attendance",
                                                      "operator": "is_in", "value": [a]}])]
    out = []
    for name, cf in raw.items():
        state = cf.get("cf_yatra_state") or cf.get("cf_event_state_new_dropdown")
        rules = CRM_BASE + [{"attribute": "cf_event_type", "operator": "is_in", "value": [name]}]
        ev = {
            "name": name,
            "type": cf.get("cf_event_type_A"),
            "date": (cf.get("cf_yatra_date") or "")[:10] or None,
            "city": cf.get("cf_yatra_city"),
            "districtField": cf.get("cf_event_district"),
            "stateText": cf.get("cf_yatra_state"),
            "stateDropdown": cf.get("cf_event_state_new_dropdown"),
            "zone": cf.get("cf_event_zone_new_dropdown"),
            "partner": cf.get("cf_organising_partner"),
            "model": partner_model(cf.get("cf_organising_partner")),
            "mode": cf.get("cf_mode_of_event"),
            "serviceLine": cf.get("cf_event_service_line1"),
            "registrations": count(rules),
            "attendance": {},
            "status": {},
            "flags": [],
        }
        if state != "Uttar Pradesh" or (ev["stateText"] and ev["stateText"] != "Uttar Pradesh"):
            ev["flags"].append("state-mismatch")
        for a in att_used:
            ev["attendance"][a] = count(rules + [{"attribute": "cf_attendance", "operator": "is_in", "value": [a]}])
        for key, (_, sid) in used.items():
            ev["status"][key] = count(rules + [{"attribute": "contact_status_id",
                                                "operator": "is_in", "value": [str(sid)]}])
        recorded = ev["attendance"].get("Attended", 0) + ev["attendance"].get("Not Attended", 0)
        if ev["registrations"] and not recorded and ev["date"] and ev["date"] < ist_today().isoformat():
            ev["flags"].append("attendance-missing")
        if ev["registrations"] < 5:
            ev["flags"].append("few-records")
        out.append(ev)
        print(f"  counted {name}", file=sys.stderr)
    return out


# --------------------------------------------------------------------------- assemble

def ist_today():
    return dt.datetime.now(IST).date()


def crm_record(e):
    return {
        "event": e["name"], "date": e["date"], "type": e["type"], "district": e["district"],
        "city": e["city"], "partner": e["partner"], "model": e["model"], "mode": e["mode"],
        "serviceLine": e["serviceLine"], "registrations": e["registrations"],
        "attendance": e["attendance"], "status": e["status"], "flags": e["flags"],
    }


def assemble(districts, plan, events):
    """Planned UP GY yatras are the core; CRM events attach to them by district + date."""
    today = ist_today()
    events = [e for e in events
              if not ("state-mismatch" in e["flags"] and e["stateText"] not in (None, "Uttar Pradesh"))]
    for e in events:
        # Prefer the city; the CRM district field is sometimes wrong (e.g. a Bihar district).
        city, field = canon(e["city"]), canon(e.get("districtField"))
        e["district"] = city if city in districts else field if field in districts else city
        if field and field != e["district"] and "district-mismatch" not in e["flags"]:
            e["flags"].append("district-mismatch")

    current = [e for e in events if (e["date"] or "") >= PROGRAMME_START]
    history = [crm_record(e) for e in events if (e["date"] or "") < PROGRAMME_START]
    used = set()
    yatras = []
    for p in plan:
        pdate = dt.date.fromisoformat(p["date"])
        best = None
        for e in current:
            if e["name"] in used or e["district"] not in p["districts"] or not e["date"]:
                continue
            gap = abs((dt.date.fromisoformat(e["date"]) - pdate).days)
            if gap <= MATCH_DAYS and (best is None or gap < best[0]):
                best = (gap, e)
        crm = None
        if best:
            used.add(best[1]["name"])
            crm = crm_record(best[1])
            crm["dayShift"] = (dt.date.fromisoformat(best[1]["date"]) - pdate).days
        held_on = dt.date.fromisoformat(crm["date"]) if crm else pdate
        yatras.append({**p, "crm": crm,
                       "state": "completed" if held_on < today else "today" if held_on == today else "upcoming"})
    unplanned = [crm_record(e) for e in current if e["name"] not in used]
    history.sort(key=lambda e: e["date"] or "")
    names = {d for y in yatras for d in y["districts"]} | {e["district"] for e in history + unplanned}
    unknown = sorted(d for d in names if d not in districts)
    return yatras, history, unplanned, unknown


def write_index(payload):
    with open(INDEX) as f:
        html = f.read()
    blob = json.dumps(payload, separators=(",", ":")).replace("</", "<\\/")
    new, n = re.subn(r"(<script id=\"portal-data\" type=\"application/json\">).*?(</script>)",
                     lambda m: m.group(1) + blob + m.group(2), html, flags=re.S)
    if n != 1:
        sys.exit("portal/index.html is missing the portal-data script block")
    with open(INDEX, "w") as f:
        f.write(new)


def main():
    fw.load_env_file()
    districts = load_districts()
    route = load_route()
    covered = [d for y in route["yatras"] for d in y["districts"]]
    if sorted(set(covered)) != sorted(districts) or len(covered) != len(set(covered)):
        sys.exit(f"Route plan districts do not match the 75 districts: "
                 f"missing {sorted(set(districts) - set(covered))}, "
                 f"unknown {sorted(set(covered) - set(districts))}")
    if "--no-crm" in sys.argv:
        events = []
    elif "--cached" in sys.argv:
        with open(CACHE) as f:
            events = json.load(f)
    else:
        events = crm_events()
        os.makedirs(os.path.dirname(CACHE), exist_ok=True)
        with open(CACHE, "w") as f:
            json.dump(events, f)
    yatras, history, unplanned, unknown = assemble(districts, route["yatras"], events)
    with open(GEOJSON) as f:
        geo = json.load(f)
    for feat in geo["features"]:
        feat["properties"] = {"name": canon(feat["properties"]["dtname"])}
    missing = sorted(set(districts) - {f["properties"]["name"] for f in geo["features"]})
    payload = {
        "generatedAt": dt.datetime.now(IST).isoformat(timespec="minutes"),
        "today": ist_today().isoformat(),
        "crmLive": bool(events),
        "programmeStart": PROGRAMME_START,
        "circularDate": CIRCULAR_DATE,
        "districts": list(districts.values()),
        "yatras": yatras,
        "reserves": route["reserves"],
        "teams": route["teams"],
        "holidays": route["holidays"],
        "fridays": route["fridays"],
        "rhythm": route["rhythm"],
        "risks": route["risks"],
        "rules": route["rules"],
        "history": history,
        "unplanned": unplanned,
        "geo": geo,
        "notes": {"unmatchedDistricts": unknown, "missingFromMap": missing},
    }
    write_index(payload)
    print(f"Wrote {len(yatras)} planned yatras ({sum(1 for y in yatras if y['crm'])} with CRM data), "
          f"{len(history)} past events, {len(unplanned)} unplanned; "
          f"{budget.total} CRM calls. Unmatched: {unknown}. Missing from map: {missing}",
          file=sys.stderr)


if __name__ == "__main__":
    main()
